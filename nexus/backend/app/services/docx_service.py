"""Geração de relatório em DOCX (RF11)."""

import io
import re

from docx import Document
from docx.shared import Pt, RGBColor


def _add_heading(doc: Document, text: str, level: int = 1) -> None:
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.color.rgb = RGBColor(0x1A, 0x56, 0xA4)


def _render_markdown_line(doc: Document, line: str) -> None:
    """Renderização simples de linhas markdown (#, ##, -, texto puro)."""
    stripped = line.rstrip()
    if not stripped:
        doc.add_paragraph("")
        return

    if stripped.startswith("### "):
        _add_heading(doc, stripped[4:], level=3)
    elif stripped.startswith("## "):
        _add_heading(doc, stripped[3:], level=2)
    elif stripped.startswith("# "):
        _add_heading(doc, stripped[2:], level=1)
    elif stripped.startswith(("- ", "* ")):
        doc.add_paragraph(stripped[2:], style="List Bullet")
    else:
        p = doc.add_paragraph()
        # Remove **bold** marks — docx-python não renderiza markdown inline
        plain = re.sub(r"\*\*(.+?)\*\*", r"\1", stripped)
        plain = re.sub(r"\*(.+?)\*", r"\1", plain)
        run = p.add_run(plain)
        run.font.size = Pt(11)


def build_report_docx(
    comparison_title: str,
    report_markdown: str,
    novelties: list,
    suggestions: list,
) -> bytes:
    doc = Document()

    title = doc.add_heading("Nexus — Relatório de Comparação", level=0)
    for run in title.runs:
        run.font.color.rgb = RGBColor(0x1A, 0x56, 0xA4)

    doc.add_paragraph(comparison_title)
    doc.add_paragraph("")

    _add_heading(doc, "Relatório da Análise", level=1)
    for line in (report_markdown or "").splitlines():
        _render_markdown_line(doc, line)

    if novelties:
        _add_heading(doc, "Novidades identificadas", level=1)
        for item in novelties:
            text = item if isinstance(item, str) else (
                item.get("text") or str(item)
            )
            doc.add_paragraph(text, style="List Bullet")

    if suggestions:
        _add_heading(doc, "Sugestões de melhoria", level=1)
        for item in suggestions:
            if isinstance(item, dict):
                text = item.get("text", "")
                score = item.get("confidence_score")
                if score is not None:
                    text = f"{text}  (confiança: {float(score):.0%})"
            else:
                text = str(item)
            doc.add_paragraph(text, style="List Bullet")

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.read()
